"""Tests for the external benchmark evaluator (rules mirror SpreadsheetBench V1)."""

import datetime
import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from spreadsheet_agents.benchmark.evaluator import (
    EvaluationResult,
    RecalculationUnavailable,
    compare_workbooks,
    evaluate_output,
    find_soffice,
    normalize_value,
    parse_answer_position,
    recalculate,
    score_run,
    values_match,
)
from spreadsheet_agents.benchmark.spreadsheetbench_v1 import (
    find_answer_position,
    load_metadata,
)

REAL_DATASET = Path("data/raw/spreadsheetbench_v1/sample_data_200")
needs_libreoffice = pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")


def save(path: Path, cells: dict, sheet: str = "Sheet1", extra_sheets: tuple = ()) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for name, value in cells.items():
        ws[name] = value
    for extra in extra_sheets:
        wb.create_sheet(extra)
    wb.save(path)
    return path


# --- value rules -------------------------------------------------------------


@pytest.mark.parametrize(
    "expected, actual",
    [
        (1, 1.0),
        (1.004, 1.001),  # both round to 1.0
        (2.5, "2.5"),  # numeric strings become numbers
        (5, " 5 "),
        (True, 1),
        (False, 0),
        (None, ""),
        ("", None),
        ("", ""),
        (None, None),
        ("abc", "abc"),
        (datetime.datetime(2021, 1, 28, 0, 0), datetime.datetime(2021, 1, 28, 3, 0)),
        (datetime.time(9, 30, 15), datetime.time(9, 30, 45)),  # seconds are dropped
    ],
)
def test_values_that_the_benchmark_treats_as_equal(expected, actual):
    assert values_match(expected, actual)


@pytest.mark.parametrize(
    "expected, actual",
    [
        (1.0, 1.02),
        (None, 0),
        (0, None),
        ("", 0),
        ("Abc", "abc"),
        ("x", "x "),
        ("abc", 5),
        (datetime.datetime(2021, 1, 28), datetime.datetime(2021, 1, 29)),
        (datetime.time(9, 30), datetime.time(9, 31)),
    ],
)
def test_values_that_the_benchmark_treats_as_different(expected, actual):
    assert not values_match(expected, actual)


def test_normalize_rounds_to_two_decimals():
    assert normalize_value(3.14159) == 3.14
    assert normalize_value("3.14159") == 3.14


# --- answer_position ---------------------------------------------------------


def test_single_cell():
    assert parse_answer_position("E1", "S") == [("S", ["E1"])]


def test_range_expands_every_cell():
    (sheet, cells), = parse_answer_position("H3:H5", "S")
    assert (sheet, cells) == ("S", ["H3", "H4", "H5"])
    (_, block), = parse_answer_position("A1:B2", "S")
    assert sorted(block) == ["A1", "A2", "B1", "B2"]


def test_multi_letter_columns():
    (_, cells), = parse_answer_position("AB1:AC1", "S")
    assert cells == ["AB1", "AC1"]


def test_quoted_sheet_name_with_spaces():
    assert parse_answer_position("'My Sheet'!A1:A2", "S") == [("My Sheet", ["A1", "A2"])]


def test_several_ranges_with_and_without_spaces():
    parsed = parse_answer_position("B12:B13, C12:C13,'Data'!D1", "S")
    assert [sheet for sheet, _ in parsed] == ["S", "S", "Data"]
    assert parsed[1][1] == ["C12", "C13"]


# --- comparing workbooks -----------------------------------------------------


def test_identical_values_pass(tmp_path):
    cells = {"H3": "D", "H4": "A"}
    answer = save(tmp_path / "a.xlsx", cells)
    output = save(tmp_path / "o.xlsx", cells)
    result = compare_workbooks(answer, output, "H3:H4")
    assert result.passed
    assert result.mismatches == []


def test_mismatch_reports_cell_expected_and_actual(tmp_path):
    answer = save(tmp_path / "a.xlsx", {"B6": 1000, "C6": 1008.33})
    output = save(tmp_path / "o.xlsx", {"B6": 1008.33, "C6": 1008.33})
    result = compare_workbooks(answer, output, "B6:C6")
    assert not result.passed
    assert [(m.cell, m.expected, m.actual) for m in result.mismatches] == [("B6", 1000, 1008.33)]


def test_only_the_answer_range_is_compared(tmp_path):
    answer = save(tmp_path / "a.xlsx", {"A1": "x", "B1": "right"})
    output = save(tmp_path / "o.xlsx", {"A1": "DIFFERENT", "B1": "right"})
    assert compare_workbooks(answer, output, "B1").passed


def test_formula_text_is_never_compared(tmp_path):
    # An output holding a formula without a cached value reads as empty: it must fail.
    answer = save(tmp_path / "a.xlsx", {"D2": 6})
    output = save(tmp_path / "o.xlsx", {"D2": "=B2*C2"})
    result = compare_workbooks(answer, output, "D2")
    assert not result.passed
    assert result.mismatches[0].actual is None


def test_missing_output_file_fails(tmp_path):
    result = compare_workbooks(save(tmp_path / "a.xlsx", {}), tmp_path / "nope.xlsx", "A1")
    assert not result.passed
    assert "does not exist" in result.error


def test_missing_sheet_in_output_fails(tmp_path):
    answer = save(tmp_path / "a.xlsx", {"A1": 1}, sheet="Data")
    output = save(tmp_path / "o.xlsx", {"A1": 1}, sheet="Other")
    result = compare_workbooks(answer, output, "'Data'!A1")
    assert not result.passed
    assert "worksheet not found" in result.error


def test_unspecified_sheet_means_the_answers_first_sheet(tmp_path):
    answer = save(tmp_path / "a.xlsx", {"A1": 1}, sheet="First", extra_sheets=("Second",))
    output = save(tmp_path / "o.xlsx", {"A1": 1}, sheet="First", extra_sheets=("Second",))
    assert compare_workbooks(answer, output, "A1").passed


def test_default_sheet_name_with_trailing_space_is_kept(tmp_path):
    # Real case: task 54590's first sheet is named "HR " (with a space).
    book = save(tmp_path / "a.xlsx", {"A1": 1}, sheet="HR ")
    assert compare_workbooks(book, book, "A1").passed


def test_every_range_of_a_multi_range_position_must_match(tmp_path):
    answer = save(tmp_path / "a.xlsx", {"A1": 1, "C1": 3})
    output = save(tmp_path / "o.xlsx", {"A1": 1, "C1": 99})
    assert not compare_workbooks(answer, output, "A1, C1").passed


def test_unreadable_output_fails_cleanly(tmp_path):
    broken = tmp_path / "broken.xlsx"
    broken.write_text("not a workbook")
    result = compare_workbooks(save(tmp_path / "a.xlsx", {}), broken, "A1")
    assert not result.passed
    assert "cannot open" in result.error


def test_bad_answer_position_fails_cleanly(tmp_path):
    book = save(tmp_path / "a.xlsx", {})
    result = compare_workbooks(book, book, "not a range")
    assert not result.passed
    assert "bad answer_position" in result.error


# --- LibreOffice recalculation (skipped where it is not installed) -----------


@needs_libreoffice
def test_recalculation_turns_formulas_into_values(tmp_path):
    output = save(tmp_path / "o.xlsx", {"B2": 2, "C2": 3, "D2": "=B2*C2"})
    answer = save(tmp_path / "a.xlsx", {"D2": 6})

    assert not evaluate_output(output, answer, "D2", recalculate_formulas=False).passed
    assert evaluate_output(output, answer, "D2").passed


@needs_libreoffice
def test_recalculation_detects_a_wrong_formula(tmp_path):
    output = save(tmp_path / "o.xlsx", {"B2": 2, "C2": 3, "D2": "=B2+C2"})
    answer = save(tmp_path / "a.xlsx", {"D2": 6})
    result = evaluate_output(output, answer, "D2")
    assert not result.passed
    assert (result.mismatches[0].expected, result.mismatches[0].actual) == (6, 5)


@needs_libreoffice
def test_recalculation_does_not_modify_the_input(tmp_path):
    output = save(tmp_path / "o.xlsx", {"B2": 2, "D2": "=B2*2"})
    before = output.read_bytes()
    recalculate(output, tmp_path / "copy.xlsx")
    assert output.read_bytes() == before


def test_missing_libreoffice_is_a_clear_error(tmp_path, monkeypatch):
    monkeypatch.setattr("spreadsheet_agents.benchmark.evaluator.find_soffice", lambda: None)
    book = save(tmp_path / "o.xlsx", {"A1": 1})
    with pytest.raises(RecalculationUnavailable, match="LibreOffice not found"):
        recalculate(book, tmp_path / "out.xlsx")


def test_evaluate_output_missing_file(tmp_path):
    result = evaluate_output(tmp_path / "nope.xlsx", save(tmp_path / "a.xlsx", {}), "A1")
    assert isinstance(result, EvaluationResult)
    assert not result.passed


# --- scoring saved runs ------------------------------------------------------


@pytest.fixture
def bench(tmp_path: Path) -> Path:
    root = tmp_path / "bench"
    task_dir = root / "spreadsheet" / "101"
    task_dir.mkdir(parents=True)
    (root / "dataset.json").write_text(
        json.dumps([{"id": 101, "instruction": "Sum", "answer_position": "B2"}])
    )
    save(task_dir / "1_101_input.xlsx", {"A2": 1})
    save(task_dir / "1_101_answer.xlsx", {"B2": 42})
    return root


def write_metrics(tmp_path: Path, **fields) -> Path:
    metrics = {"run_id": "run1", "task_id": "101", "variant": 1, "status": "success", **fields}
    path = tmp_path / "metrics.json"
    path.write_text(json.dumps(metrics))
    return path


def test_score_run_passes_a_correct_output(bench, tmp_path):
    output = save(tmp_path / "out.xlsx", {"B2": 42})
    record = score_run(
        write_metrics(tmp_path, output_path=str(output)),
        bench,
        evaluations_dir=tmp_path / "ev",
        recalculate_formulas=False,
    )
    assert record["passed"] is True
    assert json.loads((tmp_path / "ev" / "run1.json").read_text())["passed"] is True


def test_score_run_fails_a_wrong_output_and_says_which_cell(bench, tmp_path):
    output = save(tmp_path / "out.xlsx", {"B2": 41})
    record = score_run(
        write_metrics(tmp_path, output_path=str(output)),
        bench,
        evaluations_dir=tmp_path / "ev",
        recalculate_formulas=False,
    )
    assert record["passed"] is False
    assert record["mismatches"][0]["cell"] == "B2"


def test_score_run_does_not_edit_the_metrics_file(bench, tmp_path):
    output = save(tmp_path / "out.xlsx", {"B2": 42})
    metrics = write_metrics(tmp_path, output_path=str(output))
    before = metrics.read_text()
    score_run(metrics, bench, evaluations_dir=tmp_path / "ev", recalculate_formulas=False)
    assert metrics.read_text() == before


def test_run_without_output_is_a_failure_not_a_crash(bench, tmp_path):
    record = score_run(
        write_metrics(tmp_path, output_path=None, status="error"),
        bench,
        evaluations_dir=tmp_path / "ev",
    )
    assert record["passed"] is False
    assert "no output workbook" in record["error"]


def test_custom_task_cannot_be_scored(bench, tmp_path):
    record = score_run(
        write_metrics(tmp_path, task_id=None, variant=None),
        bench,
        evaluations_dir=tmp_path / "ev",
    )
    assert record["passed"] is False
    assert "not a benchmark task" in record["error"]


# --- real dataset (skipped in CI, where data/ is not committed) -------------


@pytest.mark.skipif(not REAL_DATASET.exists(), reason="SpreadsheetBench not downloaded")
def test_every_real_answer_position_parses():
    for record in load_metadata(REAL_DATASET):
        assert parse_answer_position(record["answer_position"], "Sheet1"), record["id"]


@pytest.mark.skipif(not REAL_DATASET.exists(), reason="SpreadsheetBench not downloaded")
def test_a_reference_answer_always_matches_itself():
    # Guards the whole pipeline (parser, sheet lookup, value rules) against real files.
    for record in load_metadata(REAL_DATASET)[:25]:
        task_id = record["id"]
        answer = REAL_DATASET / "spreadsheet" / str(task_id) / f"1_{task_id}_answer.xlsx"
        position = find_answer_position(REAL_DATASET, task_id)
        assert compare_workbooks(answer, answer, position).passed, task_id
