"""Tests for the LangGraph workflow: routing, state flow and failure handling."""

from pathlib import Path

import pytest
from langgraph.graph import END
from openpyxl import Workbook, load_workbook

from spreadsheet_agents.agents import editor, inspector, orchestrator, validator
from spreadsheet_agents.agents.inspector import InspectionReport
from spreadsheet_agents.graph import builder
from spreadsheet_agents.graph.builder import build_graph
from spreadsheet_agents.graph.routing import route_after_validation
from spreadsheet_agents.schemas.task import (
    EditPlan,
    SpreadsheetOperation,
    TaskPlan,
    TaskStep,
    ValidationReport,
)

# --- routing ----------------------------------------------------------------


def test_route_ends_when_validation_passed():
    assert route_after_validation({"validation_report": {"passed": True}}) == END


def test_route_fails_when_validation_did_not_pass():
    assert route_after_validation({"validation_report": {"passed": False}}) == "failed"


def test_route_fails_when_there_is_no_report():
    assert route_after_validation({}) == "failed"


# --- graph wiring (agents replaced by fakes) --------------------------------


@pytest.fixture
def initial_state(tmp_path: Path) -> dict:
    return {
        "user_request": "add a Total column",
        "input_path": str(tmp_path / "in.xlsx"),
        "output_path": str(tmp_path / "out.xlsx"),
        "model": "test-model",
    }


@pytest.fixture
def fake_agents(monkeypatch):
    """Replace the four agent entry points; record what each one receives."""
    calls: dict[str, dict] = {}
    validation = {"report": ValidationReport(passed=True)}

    def fake_inspector(workbook_path, user_request, model):
        from spreadsheet_agents.schemas.workbook import SheetInfo, WorkbookInfo

        calls["inspector"] = {"path": workbook_path}
        info = WorkbookInfo(path=workbook_path, sheets=[SheetInfo(name="Sales", rows=2, columns=2)])
        return info, InspectionReport(relevant_sheets=["Sales"])

    def fake_orchestrator(user_request, workbook_info, inspection_report, model):
        calls["orchestrator"] = {"sheets": workbook_info.sheet_names, "report": inspection_report}
        return TaskPlan(
            reasoning_summary="plan",
            steps=[TaskStep(id="s1", description="d", operation_type="set_cell")],
        )

    def fake_editor(**kwargs):
        calls["editor"] = kwargs
        return EditPlan(
            explanation="edit",
            operations=[SpreadsheetOperation(type="set_cell", sheet="Sales", cell="A1", value=1)],
        )

    def fake_validator(**kwargs):
        calls["validator"] = kwargs
        return validation["report"]

    monkeypatch.setattr(builder, "run_inspector", fake_inspector)
    monkeypatch.setattr(builder, "run_orchestrator", fake_orchestrator)
    monkeypatch.setattr(builder, "run_editor", fake_editor)
    monkeypatch.setattr(builder, "run_validator", fake_validator)
    return calls, validation


def test_successful_run_visits_agents_in_order(fake_agents, initial_state):
    final = build_graph().invoke(initial_state)

    assert [entry["agent"] for entry in final["trace"]] == [
        "inspector",
        "orchestrator",
        "editor",
        "validator",
    ]
    assert all(entry["status"] == "completed" for entry in final["trace"])
    assert final["status"] == "success"
    assert final["error"] is None


def test_each_agent_receives_the_previous_agents_output(fake_agents, initial_state):
    calls, _ = fake_agents

    build_graph().invoke(initial_state)

    assert calls["orchestrator"]["sheets"] == ["Sales"]
    assert calls["orchestrator"]["report"].relevant_sheets == ["Sales"]
    assert calls["editor"]["task_plan"].reasoning_summary == "plan"
    assert calls["editor"]["input_path"] == initial_state["input_path"]
    assert calls["editor"]["output_path"] == initial_state["output_path"]
    assert calls["validator"]["edit_plan"].explanation == "edit"
    assert calls["validator"]["output_path"] == initial_state["output_path"]


def test_final_state_keeps_every_agent_output(fake_agents, initial_state):
    final = build_graph().invoke(initial_state)

    keys = ("workbook_info", "inspection_report", "task_plan", "edit_plan", "validation_report")
    for key in keys:
        assert isinstance(final[key], dict), key


def test_failed_validation_ends_in_failed_status(fake_agents, initial_state):
    _, validation = fake_agents
    validation["report"] = ValidationReport(passed=False, errors=["A1 wrong", "B2 missing"])

    final = build_graph().invoke(initial_state)

    assert final["status"] == "failed"
    assert final["error"] == "A1 wrong; B2 missing"


def test_failed_validation_without_errors_gets_a_default_message(fake_agents, initial_state):
    _, validation = fake_agents
    validation["report"] = ValidationReport(passed=False)

    final = build_graph().invoke(initial_state)

    assert final["status"] == "failed"
    assert final["error"] == "Validation failed"


def test_agent_exception_is_not_reported_as_success(fake_agents, initial_state, monkeypatch):
    def boom(**kwargs):
        raise RuntimeError("editor exploded")

    monkeypatch.setattr(builder, "run_editor", boom)

    with pytest.raises(RuntimeError, match="editor exploded"):
        build_graph().invoke(initial_state)


# --- end to end (real tools, scripted fake LLM) -----------------------------


class ScriptedLLM:
    """Answers each agent with a canned structured result, chosen by output schema."""

    answers: dict = {}

    def __init__(self, model: str, temperature: float):
        self.schema = None

    def with_structured_output(self, schema):
        self.schema = schema
        return self

    def invoke(self, messages):
        return self.answers[self.schema]


def test_end_to_end_with_real_tools(tmp_path, monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Units", "Price"])
    ws.append(["North", 10, 2.5])
    source = tmp_path / "in.xlsx"
    wb.save(source)
    original = source.read_bytes()
    out = tmp_path / "out" / "result.xlsx"

    ScriptedLLM.answers = {
        InspectionReport: InspectionReport(relevant_sheets=["Sales"]),
        TaskPlan: TaskPlan(reasoning_summary="multiply"),
        EditPlan: EditPlan(
            explanation="add total",
            operations=[
                SpreadsheetOperation(type="set_cell", sheet="Sales", cell="D1", value="Total"),
                SpreadsheetOperation(
                    type="set_formula", sheet="Sales", cell="D2", value="=B2*C2"
                ),
            ],
        ),
        ValidationReport: ValidationReport(passed=True),
    }
    for module in (inspector, orchestrator, editor, validator):
        monkeypatch.setattr(module, "ChatOpenAI", ScriptedLLM)

    final = build_graph().invoke(
        {
            "user_request": "add a Total column",
            "input_path": str(source),
            "output_path": str(out),
            "model": "test-model",
        }
    )

    assert final["status"] == "success"
    result = load_workbook(out)["Sales"]
    assert result["D1"].value == "Total"
    assert result["D2"].value == "=B2*C2"
    assert source.read_bytes() == original
    assert final["validation_report"]["checked_operations"] == 2


def test_end_to_end_reports_failure_when_validator_finds_a_mismatch(tmp_path, monkeypatch):
    wb = Workbook()
    wb.active.title = "Sales"
    source = tmp_path / "in.xlsx"
    wb.save(source)

    ScriptedLLM.answers = {
        InspectionReport: InspectionReport(),
        TaskPlan: TaskPlan(reasoning_summary="x"),
        EditPlan: EditPlan(
            explanation="e",
            operations=[
                SpreadsheetOperation(type="set_cell", sheet="Sales", cell="A1", value="ok")
            ],
        ),
        ValidationReport: ValidationReport(passed=False, errors=["semantic problem"]),
    }
    for module in (inspector, orchestrator, editor, validator):
        monkeypatch.setattr(module, "ChatOpenAI", ScriptedLLM)

    final = build_graph().invoke(
        {
            "user_request": "r",
            "input_path": str(source),
            "output_path": str(tmp_path / "out.xlsx"),
            "model": "m",
        }
    )

    assert final["status"] == "failed"
    assert final["error"] == "semantic problem"
