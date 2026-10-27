"""Tests for the Editor agent with a fake LLM (no network)."""

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from spreadsheet_agents.agents import editor
from spreadsheet_agents.agents.editor import generate_edit_plan, run_editor
from spreadsheet_agents.agents.inspector import InspectionReport
from spreadsheet_agents.schemas.task import (
    EditPlan,
    SpreadsheetOperation,
    TaskPlan,
    TaskStep,
)
from spreadsheet_agents.schemas.workbook import SheetInfo, WorkbookInfo


class FakeLLM:
    """Stands in for ChatOpenAI and records how the agent called it."""

    instances: list["FakeLLM"] = []
    response = None  # set per test with monkeypatch

    def __init__(self, model: str, temperature: float):
        self.model = model
        self.temperature = temperature
        self.schema = None
        self.messages = None
        FakeLLM.instances.append(self)

    def with_structured_output(self, schema, **kwargs):
        self.schema = schema
        self.kwargs = kwargs
        return self

    def invoke(self, messages):
        self.messages = messages
        return self.response


@pytest.fixture
def fake_llm(monkeypatch):
    FakeLLM.instances = []
    monkeypatch.setattr(editor, "ChatOpenAI", FakeLLM)
    return FakeLLM


@pytest.fixture
def source(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Units", "Price"])
    ws.append(["North", 10, 2.5])
    path = tmp_path / "in.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def context():
    return {
        "user_request": "add a Total column",
        "task_plan": TaskPlan(
            reasoning_summary="multiply units by price",
            steps=[TaskStep(id="s1", description="add total", operation_type="set_formula")],
        ),
        "workbook_info": WorkbookInfo(
            path="in.xlsx", sheets=[SheetInfo(name="Sales", rows=2, columns=3)]
        ),
        "inspection_report": InspectionReport(observations=["Units and Price found"]),
    }


def total_plan() -> EditPlan:
    return EditPlan(
        explanation="Add header and formula",
        operations=[
            SpreadsheetOperation(type="set_cell", sheet="Sales", cell="D1", value="Total"),
            SpreadsheetOperation(type="set_formula", sheet="Sales", cell="D2", value="=B2*C2"),
        ],
    )


def test_generate_returns_structured_plan(fake_llm, context, monkeypatch):
    plan = total_plan()
    monkeypatch.setattr(FakeLLM, "response", plan)

    result = generate_edit_plan(**context, model="test-model")

    assert result == plan
    llm = fake_llm.instances[0]
    assert llm.schema is EditPlan
    assert llm.kwargs == {"method": "function_calling"}
    assert llm.model == "test-model"
    assert llm.temperature == 0


def test_generate_sends_all_context_to_the_model(fake_llm, context, monkeypatch):
    monkeypatch.setattr(FakeLLM, "response", total_plan())

    generate_edit_plan(**context, model="test-model")

    llm = fake_llm.instances[0]
    text = "\n".join(content for _, content in llm.messages)
    assert "add a Total column" in text
    assert "multiply units by price" in text
    assert "Sales" in text
    assert "Units and Price found" in text
    assert "Editing Agent" in llm.messages[0][1]


def test_generate_rejects_empty_plan(fake_llm, context, monkeypatch):
    monkeypatch.setattr(FakeLLM, "response", EditPlan(explanation="cannot be done"))

    with pytest.raises(ValueError, match="no operations.*cannot be done"):
        generate_edit_plan(**context, model="test-model")


def test_run_editor_executes_plan_on_a_new_copy(fake_llm, context, source, tmp_path, monkeypatch):
    monkeypatch.setattr(FakeLLM, "response", total_plan())
    out = tmp_path / "out.xlsx"
    before = source.read_bytes()

    plan = run_editor(**context, input_path=str(source), output_path=str(out), model="m")

    sales = load_workbook(out)["Sales"]
    assert sales["D1"].value == "Total"
    assert sales["D2"].value == "=B2*C2"
    assert plan == total_plan()
    assert source.read_bytes() == before


def test_run_editor_writes_nothing_when_plan_is_empty(
    fake_llm, context, source, tmp_path, monkeypatch
):
    monkeypatch.setattr(FakeLLM, "response", EditPlan(explanation="nothing to do"))
    out = tmp_path / "out.xlsx"

    with pytest.raises(ValueError):
        run_editor(**context, input_path=str(source), output_path=str(out), model="m")

    assert not out.exists()


def test_run_editor_surfaces_invalid_llm_operations(
    fake_llm, context, source, tmp_path, monkeypatch
):
    bad = EditPlan(
        explanation="bad",
        operations=[SpreadsheetOperation(type="set_cell", sheet="Ghost", cell="A1", value=1)],
    )
    monkeypatch.setattr(FakeLLM, "response", bad)
    out = tmp_path / "out.xlsx"

    with pytest.raises(RuntimeError, match="Operation 0 failed.*Unknown worksheet"):
        run_editor(**context, input_path=str(source), output_path=str(out), model="m")

    assert not out.exists()
