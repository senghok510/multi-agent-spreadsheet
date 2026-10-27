"""Tests for deterministic workbook inspection (no LLM involved)."""

from pathlib import Path

import pytest
from openpyxl import Workbook

from spreadsheet_agents.tools.workbook_reader import inspect_workbook


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    wb = Workbook()
    sales = wb.active
    sales.title = "Sales"
    sales.append(["Region", "Units", "Price", "Total"])
    for i, region in enumerate(["North", "South", "East"], start=2):
        sales.append([region, i * 10, 2.5, f"=B{i}*C{i}"])

    notes = wb.create_sheet("Notes")
    notes["A1"] = "Remember to check the totals"

    wb.create_sheet("Empty")

    path = tmp_path / "sample.xlsx"
    wb.save(path)
    return path


def test_reports_sheet_names_in_order(workbook_path):
    info = inspect_workbook(workbook_path)
    assert info.sheet_names == ["Sales", "Notes", "Empty"]


def test_reports_dimensions_and_headers(workbook_path):
    sales = inspect_workbook(workbook_path).sheets[0]
    assert (sales.rows, sales.columns) == (4, 4)
    assert sales.headers == ["Region", "Units", "Price", "Total"]


def test_sample_rows_respect_limit(workbook_path):
    sales = inspect_workbook(workbook_path, max_preview_rows=2).sheets[0]
    assert len(sales.sample_rows) == 2
    assert sales.sample_rows[1][0] == "North"


def test_formulas_are_reported_not_evaluated(workbook_path):
    sales = inspect_workbook(workbook_path).sheets[0]
    assert {"cell": "D2", "formula": "=B2*C2"} in sales.formulas
    assert len(sales.formulas) == 3


def test_formula_count_is_capped(workbook_path):
    sales = inspect_workbook(workbook_path, max_formulas=1).sheets[0]
    assert len(sales.formulas) == 1


def test_empty_sheet_does_not_crash(workbook_path):
    empty = inspect_workbook(workbook_path).sheets[2]
    assert empty.formulas == []


def test_input_file_is_not_modified(workbook_path):
    before = workbook_path.read_bytes()
    inspect_workbook(workbook_path)
    assert workbook_path.read_bytes() == before


def test_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        inspect_workbook(tmp_path / "nope.xlsx")
