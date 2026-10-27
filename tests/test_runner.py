"""Tests for the task runner: output files, metrics and error handling (no network)."""

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from spreadsheet_agents.agents import editor, inspector, orchestrator, validator
from spreadsheet_agents.agents.inspector import InspectionReport
from spreadsheet_agents.runner import run_task
from spreadsheet_agents.schemas.task import (
    EditPlan,
    SpreadsheetOperation,
    TaskPlan,
    ValidationReport,
)


class ScriptedLLM:
    answers: dict = {}

    def __init__(self, model: str, temperature: float):
        self.schema = None

    def with_structured_output(self, schema, **kwargs):
        self.schema = schema
        self.kwargs = kwargs
        return self

    def invoke(self, messages):
        return self.answers[self.schema]


@pytest.fixture
def source(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Units"])
    path = tmp_path / "task.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def scripted(monkeypatch):
    ScriptedLLM.answers = {
        InspectionReport: InspectionReport(relevant_sheets=["Sales"]),
        TaskPlan: TaskPlan(reasoning_summary="write total"),
        EditPlan: EditPlan(
            explanation="e",
            operations=[SpreadsheetOperation(type="set_cell", sheet="Sales", cell="C1", value="T")],
        ),
        ValidationReport: ValidationReport(passed=True),
    }
    for module in (inspector, orchestrator, editor, validator):
        monkeypatch.setattr(module, "ChatOpenAI", ScriptedLLM)
    return ScriptedLLM


def run(source, tmp_path, **extra):
    return run_task(
        instruction="write T in C1",
        input_path=source,
        model="test-model",
        output_dir=tmp_path / "outputs",
        **extra,
    )


def test_successful_run_writes_workbook_trace_and_metrics(scripted, source, tmp_path):
    metrics = run(source, tmp_path, label="demo")

    assert metrics["status"] == "success"
    assert load_workbook(metrics["output_path"])["Sales"]["C1"].value == "T"
    assert Path(metrics["trace_path"]).exists()
    assert Path(metrics["metrics_path"]).exists()
    assert metrics["run_id"].endswith("_demo")


def test_input_workbook_is_not_modified(scripted, source, tmp_path):
    before = source.read_bytes()
    run(source, tmp_path)
    assert source.read_bytes() == before


def test_metrics_record_the_run_and_stay_honest_about_evaluation(scripted, source, tmp_path):
    metrics = run(source, tmp_path, task_id="59196", variant=2)

    assert metrics["model"] == "test-model"
    assert metrics["task_id"] == "59196"
    assert metrics["variant"] == 2
    assert metrics["operations_planned"] == 1
    assert metrics["latency_seconds"] >= 0
    assert metrics["runtime_validation"]["passed"] is True
    assert metrics["benchmark_evaluated"] is False


def test_metrics_file_matches_returned_metrics(scripted, source, tmp_path):
    metrics = run(source, tmp_path)
    on_disk = json.loads(Path(metrics["metrics_path"]).read_text())
    assert on_disk == json.loads(json.dumps(metrics, default=str))


def test_trace_file_lists_the_four_agents_in_order(scripted, source, tmp_path):
    metrics = run(source, tmp_path)
    trace = json.loads(Path(metrics["trace_path"]).read_text())

    assert [step["agent"] for step in trace["steps"]] == [
        "inspector",
        "orchestrator",
        "editor",
        "validator",
    ]
    assert trace["instruction"] == "write T in C1"
    assert trace["edit_plan"]["operations"][0]["cell"] == "C1"


def test_failed_validation_is_recorded_as_failed(scripted, source, tmp_path):
    scripted.answers[ValidationReport] = ValidationReport(passed=False, errors=["wrong"])

    metrics = run(source, tmp_path)

    assert metrics["status"] == "failed"
    assert metrics["error"] == "wrong"
    assert metrics["output_path"] is not None  # the attempt is kept for inspection


def test_crash_is_recorded_instead_of_raised(scripted, source, tmp_path, monkeypatch):
    class Exploding(ScriptedLLM):
        def invoke(self, messages):
            raise ConnectionError("api down")

    monkeypatch.setattr(inspector, "ChatOpenAI", Exploding)

    metrics = run(source, tmp_path)

    assert metrics["status"] == "error"
    assert "ConnectionError: api down" in metrics["error"]
    assert metrics["output_path"] is None
    assert Path(metrics["metrics_path"]).exists()
    assert Path(metrics["trace_path"]).exists()
