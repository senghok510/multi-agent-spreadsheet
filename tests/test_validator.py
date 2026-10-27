"""Tests for the Validator agent: deterministic checks plus a fake-LLM semantic pass."""

from pathlib import Path

import pytest
from openpyxl import Workbook

from spreadsheet_agents.agents import validator
from spreadsheet_agents.agents.validator import deterministic_validate, run_validator
from spreadsheet_agents.schemas.task import (
    EditPlan,
    SpreadsheetOperation,
    TaskPlan,
    ValidationReport,
)


def op(type_: str, **fields) -> SpreadsheetOperation:
    return SpreadsheetOperation(type=type_, **fields)


def plan(*operations: SpreadsheetOperation) -> EditPlan:
    return EditPlan(explanation="test", operations=list(operations))


@pytest.fixture
def output(tmp_path: Path) -> Path:
    """A workbook as the editor would have left it."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "Region"
    ws["D1"] = "Total"
    ws["D2"] = "=B2*C2"
    wb.create_sheet("Summary")
    path = tmp_path / "out.xlsx"
    wb.save(path)
    return path


class FakeLLM:
    instances: list["FakeLLM"] = []
    response = None  # set per test with monkeypatch

    def __init__(self, model: str, temperature: float):
        self.messages = None
        self.schema = None
        FakeLLM.instances.append(self)

    def with_structured_output(self, schema, **kwargs):
        self.schema = schema
        self.kwargs = kwargs
        return self

    def invoke(self, messages):
        self.messages = messages
        return self.response


@pytest.fixture
def fake_llm(monkeypatch):
    FakeLLM.instances = []
    monkeypatch.setattr(validator, "ChatOpenAI", FakeLLM)
    return FakeLLM


# --- deterministic_validate -------------------------------------------------


def test_passes_when_edits_are_present(output):
    report = deterministic_validate(
        str(output),
        plan(
            op("set_cell", sheet="Sales", cell="D1", value="Total"),
            op("set_formula", sheet="Sales", cell="D2", value="=B2*C2"),
            op("create_sheet", sheet="Summary"),
        ),
    )
    assert report.passed
    assert report.checked_operations == 3
    assert report.errors == []


def test_detects_wrong_cell_value(output):
    report = deterministic_validate(
        str(output), plan(op("set_cell", sheet="Sales", cell="D1", value="Sum"))
    )
    assert not report.passed
    assert "Sales!D1" in report.errors[0]
    assert "'Sum'" in report.errors[0] and "'Total'" in report.errors[0]


def test_detects_missing_sheet(output):
    report = deterministic_validate(
        str(output), plan(op("set_cell", sheet="Ghost", cell="A1", value=1))
    )
    assert not report.passed
    assert "Missing sheet: Ghost" in report.errors


def test_detects_missing_created_sheet(output):
    report = deterministic_validate(str(output), plan(op("create_sheet", sheet="Extra")))
    assert not report.passed
    assert "Missing created sheet: Extra" in report.errors


def test_clear_cell_expects_empty(output):
    ok = deterministic_validate(str(output), plan(op("clear_cell", sheet="Sales", cell="B5")))
    bad = deterministic_validate(str(output), plan(op("clear_cell", sheet="Sales", cell="A1")))
    assert ok.passed
    assert not bad.passed


def test_overwritten_cell_is_judged_by_its_final_value(output):
    # The editor applies operations in order, so only the last write to D1 matters.
    report = deterministic_validate(
        str(output),
        plan(
            op("set_cell", sheet="Sales", cell="D1", value="Draft"),
            op("set_cell", sheet="Sales", cell="D1", value="Total"),
        ),
    )
    assert report.passed


def copy_op(source_range, destination_cell, source_sheet="Sales", destination_sheet="Sales"):
    return op(
        "copy_range",
        source_sheet=source_sheet,
        source_range=source_range,
        destination_sheet=destination_sheet,
        destination_cell=destination_cell,
    )


@pytest.fixture
def fill_down(tmp_path: Path) -> Workbook:
    """Output of: set_formula D2, copy D2 -> D3 (done correctly by the editor)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["D2"] = "=B2*C2"
    ws["D3"] = "=B3*C3"
    return wb


def save(wb: Workbook, tmp_path: Path) -> str:
    path = tmp_path / "copy.xlsx"
    wb.save(path)
    return str(path)


def test_copy_range_with_adjusted_formula_passes(fill_down, tmp_path):
    report = deterministic_validate(
        save(fill_down, tmp_path),
        plan(op("set_formula", sheet="Sales", cell="D2", value="=B2*C2"), copy_op("D2", "D3")),
    )
    assert report.passed, report.errors
    assert report.checked_operations == 2
    assert report.warnings == []


def test_copy_range_detects_formula_that_was_not_adjusted(fill_down, tmp_path):
    fill_down["Sales"]["D3"] = "=B2*C2"  # the verbatim-copy bug from the first real run
    report = deterministic_validate(save(fill_down, tmp_path), plan(copy_op("D2", "D3")))
    assert not report.passed
    assert "Sales!D3" in report.errors[0]
    assert "=B3*C3" in report.errors[0]


def test_copy_range_detects_missing_destination_value(fill_down, tmp_path):
    fill_down["Sales"]["D3"] = None
    report = deterministic_validate(save(fill_down, tmp_path), plan(copy_op("D2", "D3")))
    assert not report.passed


def test_copy_range_checks_every_cell_of_a_block(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"], ws["B1"], ws["A2"], ws["B2"] = 1, 2, 3, 4
    ws["D1"], ws["E1"], ws["D2"], ws["E2"] = 1, 2, 3, 99  # last cell is wrong
    report = deterministic_validate(save(wb, tmp_path), plan(copy_op("A1:B2", "D1")))
    assert not report.passed
    assert len(report.errors) == 1
    assert "Sales!E2" in report.errors[0]


def test_copy_range_to_another_sheet(tmp_path):
    wb = Workbook()
    wb.active.title = "Sales"
    wb["Sales"]["A1"] = "Region"
    wb.create_sheet("Summary")["A1"] = "Region"
    report = deterministic_validate(
        save(wb, tmp_path), plan(copy_op("A1", "A1", destination_sheet="Summary"))
    )
    assert report.passed


def test_copy_range_with_missing_sheet_is_an_error(tmp_path):
    wb = Workbook()
    wb.active.title = "Sales"
    report = deterministic_validate(
        save(wb, tmp_path), plan(copy_op("A1", "A1", destination_sheet="Ghost"))
    )
    assert not report.passed
    assert "Missing sheet: Ghost" in report.errors


@pytest.mark.parametrize(
    "operations, reason",
    [
        ([copy_op("A1:A3", "A2")], "overlap"),
        (
            [copy_op("D2", "D3"), op("set_cell", sheet="Sales", cell="D3", value="later")],
            "later operation",
        ),
        ([copy_op("not-a-range", "D3")], "invalid range"),
    ],
)
def test_copy_range_that_cannot_be_checked_soundly_warns_instead(
    fill_down, tmp_path, operations, reason
):
    report = deterministic_validate(save(fill_down, tmp_path), plan(*operations))
    assert any("not verified" in w and reason in w for w in report.warnings), report.warnings


def test_unreadable_output_is_a_failed_report_not_a_crash(tmp_path):
    broken = tmp_path / "broken.xlsx"
    broken.write_text("this is not a workbook")
    report = deterministic_validate(str(broken), plan(op("create_sheet", sheet="X")))
    assert not report.passed
    assert "readable" in report.errors[0].lower()


def test_missing_output_is_a_failed_report(tmp_path):
    report = deterministic_validate(
        str(tmp_path / "nope.xlsx"), plan(op("create_sheet", sheet="X"))
    )
    assert not report.passed


# --- run_validator ----------------------------------------------------------


def test_llm_is_skipped_when_deterministic_checks_fail(fake_llm, output):
    report = run_validator(
        "add total",
        TaskPlan(reasoning_summary="x"),
        plan(op("set_cell", sheet="Sales", cell="D1", value="Wrong")),
        str(output),
        model="m",
    )
    assert not report.passed
    assert fake_llm.instances == []


def test_semantic_failure_fails_the_overall_report(fake_llm, output, monkeypatch):
    monkeypatch.setattr(
        FakeLLM,
        "response",
        ValidationReport(passed=False, errors=["totals not summed"], warnings=["unsure"]),
    )
    report = run_validator(
        "add total",
        TaskPlan(reasoning_summary="x"),
        plan(op("set_cell", sheet="Sales", cell="D1", value="Total")),
        str(output),
        model="m",
    )
    assert not report.passed
    assert report.errors == ["totals not summed"]
    assert report.warnings == ["unsure"]
    assert report.checked_operations == 1


def test_both_passing_gives_passed_report(fake_llm, output, monkeypatch):
    monkeypatch.setattr(FakeLLM, "response", ValidationReport(passed=True))
    report = run_validator(
        "add total",
        TaskPlan(reasoning_summary="x"),
        plan(op("set_cell", sheet="Sales", cell="D1", value="Total")),
        str(output),
        model="m",
    )
    assert report.passed


def test_llm_sees_request_plan_and_output_preview(fake_llm, output, monkeypatch):
    monkeypatch.setattr(FakeLLM, "response", ValidationReport(passed=True))
    run_validator(
        "add total",
        TaskPlan(reasoning_summary="multiply columns"),
        plan(op("set_cell", sheet="Sales", cell="D1", value="Total")),
        str(output),
        model="m",
    )
    llm = fake_llm.instances[0]
    text = "\n".join(content for _, content in llm.messages)
    assert llm.schema is ValidationReport
    assert llm.kwargs == {"method": "function_calling"}
    assert "add total" in text
    assert "multiply columns" in text
    assert "Region" in text  # preview of the real output workbook
