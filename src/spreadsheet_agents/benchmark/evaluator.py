"""Score an output workbook against SpreadsheetBench V1's reference answer.

This is the external evaluator. It is deliberately separate from the
runtime Validator and must never be imported by agent code: it reads the
reference answer and the answer range.

The comparison follows the benchmark's own evaluation/evaluation.py:
values only (never formulas or formatting), over the cells named by
`answer_position`, after recalculating formulas.
"""

import argparse
import datetime
import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from pydantic import BaseModel, Field

from spreadsheet_agents.benchmark.spreadsheetbench_v1 import (
    find_answer_path,
    find_answer_position,
    load_task,
)
from spreadsheet_agents.schemas.task import EditPlan
from spreadsheet_agents.tools.workbook_editor import apply_operations

MAC_SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
RECALC_TIMEOUT_SECONDS = 120
VARIANTS = (1, 2, 3)  # the benchmark scores every task on three test workbooks
MAX_MISMATCHES_KEPT = 10


class RecalculationUnavailable(RuntimeError):
    """LibreOffice is needed to calculate formulas and was not found."""


class Mismatch(BaseModel):
    sheet: str
    cell: str
    expected: Any
    actual: Any


class EvaluationResult(BaseModel):
    passed: bool
    mismatches: list[Mismatch] = Field(default_factory=list)
    error: str | None = None


# --- value comparison (mirrors the benchmark) --------------------------------


def _datetime_to_serial(value: datetime.datetime) -> float:
    delta = value - datetime.datetime(1899, 12, 30)
    return delta.days + delta.seconds / 86400.0


def normalize_value(value: Any) -> Any:
    if isinstance(value, (int, float)):  # includes bool
        return round(float(value), 2)

    if isinstance(value, datetime.time):
        return str(value)[:-3]

    if isinstance(value, datetime.datetime):
        return round(_datetime_to_serial(value), 0)

    if isinstance(value, str):
        try:
            return round(float(value), 2)
        except ValueError:
            return value

    return value


def values_match(expected: Any, actual: Any) -> bool:
    expected = normalize_value(expected)
    actual = normalize_value(actual)

    if expected in ("", None) and actual in ("", None):
        return True

    if type(expected) is not type(actual):
        return False

    return expected == actual


# --- answer_position ---------------------------------------------------------


def parse_answer_position(
    answer_position: str,
    default_sheet: str,
) -> list[tuple[str, list[str]]]:
    """Split "'Sheet 1'!A1:B2, C3" into [(sheet, [cells...]), ...].

    Whitespace around each part is ignored (one real task writes
    "B12:B110, C12:C23").
    """

    parsed = []

    for part in answer_position.split(","):
        part = part.strip()

        if not part:
            continue

        if "!" in part:
            sheet, cell_range = part.rsplit("!", 1)
            sheet = sheet.strip().strip("'")
        else:
            # The workbook's real first sheet name is used as is: some are
            # padded with spaces (task 54590 has "HR ").
            sheet, cell_range = default_sheet, part

        cell_range = cell_range.strip().strip("'")

        min_col, min_row, max_col, max_row = range_boundaries(cell_range)

        parsed.append(
            (
                sheet,
                [
                    f"{get_column_letter(col)}{row}"
                    for col in range(min_col, max_col + 1)
                    for row in range(min_row, max_row + 1)
                ],
            )
        )

    return parsed


# --- comparison --------------------------------------------------------------


def compare_workbooks(
    answer_path: str | Path,
    output_path: str | Path,
    answer_position: str,
) -> EvaluationResult:
    """Compare cached cell values in the answer range. Formulas are NOT
    calculated here; use evaluate_output() for that."""

    if not Path(output_path).exists():
        return EvaluationResult(passed=False, error="output workbook does not exist")

    try:
        answer_wb = load_workbook(answer_path, data_only=True)
        output_wb = load_workbook(output_path, data_only=True)
    except Exception as exc:
        return EvaluationResult(passed=False, error=f"cannot open workbook: {exc}")

    try:
        try:
            ranges = parse_answer_position(answer_position, answer_wb.sheetnames[0])
        except Exception as exc:
            return EvaluationResult(
                passed=False, error=f"bad answer_position {answer_position!r}: {exc}"
            )

        mismatches = []

        for sheet, cells in ranges:
            if sheet not in output_wb.sheetnames:
                return EvaluationResult(
                    passed=False, error=f"worksheet not found in output: {sheet}"
                )

            if sheet not in answer_wb.sheetnames:
                return EvaluationResult(
                    passed=False, error=f"worksheet not found in answer: {sheet}"
                )

            for cell in cells:
                expected = answer_wb[sheet][cell].value
                actual = output_wb[sheet][cell].value

                if not values_match(expected, actual):
                    mismatches.append(
                        Mismatch(sheet=sheet, cell=cell, expected=expected, actual=actual)
                    )

        return EvaluationResult(passed=not mismatches, mismatches=mismatches)

    finally:
        answer_wb.close()
        output_wb.close()


# --- recalculation -----------------------------------------------------------


def find_soffice() -> str | None:
    configured = os.environ.get("SOFFICE_PATH")

    for candidate in (configured, shutil.which("soffice"), MAC_SOFFICE):
        if candidate and Path(candidate).exists():
            return candidate

    return None


def recalculate(input_path: str | Path, output_path: str | Path) -> Path:
    """Open the workbook in headless LibreOffice and save it with computed values."""

    soffice = find_soffice()

    if soffice is None:
        raise RecalculationUnavailable(
            "LibreOffice not found. Install it (brew install --cask libreoffice) "
            "or set SOFFICE_PATH."
        )

    input_path = Path(input_path).resolve()
    output_path = Path(output_path)

    with tempfile.TemporaryDirectory() as scratch:
        scratch_dir = Path(scratch)
        converted_dir = scratch_dir / "converted"

        completed = subprocess.run(
            [
                soffice,
                f"-env:UserInstallation={(scratch_dir / 'profile').as_uri()}",
                "--headless",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(converted_dir),
                str(input_path),
            ],
            capture_output=True,
            text=True,
            timeout=RECALC_TIMEOUT_SECONDS,
        )

        converted = converted_dir / f"{input_path.stem}.xlsx"

        if completed.returncode != 0 or not converted.exists():
            raise RuntimeError(
                f"LibreOffice could not convert {input_path.name}: "
                f"{completed.stderr.strip()}"
            )

        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(converted, output_path)

    return output_path


def evaluate_output(
    output_path: str | Path,
    answer_path: str | Path,
    answer_position: str,
    recalculate_formulas: bool = True,
) -> EvaluationResult:
    """Recalculate the output (unless told not to), then compare it with the answer."""

    if not Path(output_path).exists():
        return EvaluationResult(passed=False, error="output workbook does not exist")

    if not recalculate_formulas:
        return compare_workbooks(answer_path, output_path, answer_position)

    with tempfile.TemporaryDirectory() as scratch:
        recalculated = recalculate(output_path, Path(scratch) / "recalculated.xlsx")
        return compare_workbooks(answer_path, recalculated, answer_position)


# --- scoring a saved run -----------------------------------------------------


def score_run(
    metrics_path: str | Path,
    benchmark_root: str | Path,
    evaluations_dir: str | Path = "outputs/evaluations",
    recalculate_formulas: bool = True,
) -> dict[str, Any]:
    """Score one saved run and write the result next to, not inside, its metrics."""

    metrics = json.loads(Path(metrics_path).read_text(encoding="utf-8"))

    record: dict[str, Any] = {
        "run_id": metrics["run_id"],
        "task_id": metrics.get("task_id"),
        "variant": metrics.get("variant"),
        "runtime_status": metrics.get("status"),
    }

    if metrics.get("task_id") is None or metrics.get("variant") is None:
        record.update(passed=False, error="not a benchmark task; nothing to score against")
    elif not metrics.get("output_path"):
        record.update(passed=False, error="the run produced no output workbook")
    else:
        answer_path = find_answer_path(benchmark_root, metrics["task_id"], metrics["variant"])

        if answer_path is None:
            record.update(passed=False, error="no reference answer file for this task")
        else:
            result = evaluate_output(
                metrics["output_path"],
                answer_path,
                find_answer_position(benchmark_root, metrics["task_id"]),
                recalculate_formulas=recalculate_formulas,
            )
            record.update(result.model_dump())

    path = Path(evaluations_dir) / f"{metrics['run_id']}.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(record, indent=2, default=str), encoding="utf-8")
    record["evaluation_path"] = str(path)

    return record


def _load_edit_plan(metrics: dict[str, Any]) -> EditPlan | None:
    trace_path = metrics.get("trace_path")

    if not trace_path or not Path(trace_path).exists():
        return None

    plan = json.loads(Path(trace_path).read_text(encoding="utf-8")).get("edit_plan")

    return EditPlan.model_validate(plan) if plan else None


def score_run_all_variants(
    metrics_path: str | Path,
    benchmark_root: str | Path,
    evaluations_dir: str | Path = "outputs/evaluations",
    recalculate_formulas: bool = True,
) -> dict[str, Any]:
    """Score a saved run on all three benchmark variants.

    The workbook the run produced is scored for the variant it ran on. The
    saved edit plan is replayed on the other variants, which is how the
    benchmark treats a solution: one answer, run on every test case. No
    model is called. A plan that cannot be applied to a variant fails it.

    soft = fraction of variants passed; hard = 1 only if all pass.
    """

    metrics = json.loads(Path(metrics_path).read_text(encoding="utf-8"))
    task_id = metrics.get("task_id")
    own_variant = metrics.get("variant")

    record: dict[str, Any] = {
        "run_id": metrics["run_id"],
        "task_id": task_id,
        "runtime_status": metrics.get("status"),
        "variants": [],
    }

    if task_id is None or own_variant is None:
        record.update(
            passed=False,
            soft=0.0,
            hard=0,
            error="not a benchmark task; nothing to score against",
        )
    else:
        answer_position = find_answer_position(benchmark_root, task_id)
        edit_plan = _load_edit_plan(metrics)

        with tempfile.TemporaryDirectory() as scratch:
            for variant in VARIANTS:
                entry: dict[str, Any] = {
                    "variant": variant,
                    "passed": False,
                    "error": None,
                    "mismatches": [],
                }

                try:
                    answer_path = find_answer_path(benchmark_root, task_id, variant)

                    if answer_path is None:
                        entry["error"] = "no reference answer file for this variant"
                    elif variant == own_variant:
                        entry["source"] = "run output"
                        output = metrics.get("output_path")

                        if output:
                            result = evaluate_output(
                                output, answer_path, answer_position, recalculate_formulas
                            )
                        else:
                            result = EvaluationResult(
                                passed=False, error="the run produced no output workbook"
                            )

                        entry.update(_result_fields(result))
                    elif edit_plan is None:
                        entry["source"] = "replayed plan"
                        entry["error"] = "the run has no edit plan to replay"
                    else:
                        entry["source"] = "replayed plan"
                        replayed = apply_operations(
                            load_task(benchmark_root, task_id, variant).input_path,
                            Path(scratch) / f"variant{variant}.xlsx",
                            edit_plan.operations,
                        )
                        result = evaluate_output(
                            replayed, answer_path, answer_position, recalculate_formulas
                        )
                        entry.update(_result_fields(result))
                except RecalculationUnavailable:
                    raise
                except Exception as exc:
                    entry["error"] = f"{type(exc).__name__}: {exc}"

                record["variants"].append(entry)

        passed = sum(entry["passed"] for entry in record["variants"])
        record["soft"] = round(passed / len(VARIANTS), 4)
        record["hard"] = int(passed == len(VARIANTS))
        record["passed"] = bool(record["hard"])

    path = Path(evaluations_dir) / f"{metrics['run_id']}.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(record, indent=2, default=str), encoding="utf-8")
    record["evaluation_path"] = str(path)

    return record


def _result_fields(result: EvaluationResult) -> dict[str, Any]:
    fields = result.model_dump()
    fields["mismatches"] = fields["mismatches"][:MAX_MISMATCHES_KEPT]
    return fields


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Score saved runs against the SpreadsheetBench reference answers."
    )
    parser.add_argument("metrics", nargs="+", help="Path(s) to outputs/metrics/<run_id>.json")
    parser.add_argument("--benchmark-root", help="Benchmark folder (default: config)")
    parser.add_argument("--config", default="configs/baseline.yaml")
    parser.add_argument("--output-dir", default="outputs", help="Where evaluations are written")
    parser.add_argument(
        "--single-variant",
        action="store_true",
        help="Score only the variant the run used (default: all three, replaying the plan)",
    )
    args = parser.parse_args(argv)

    root = args.benchmark_root

    if root is None:
        import yaml

        root = yaml.safe_load(Path(args.config).read_text(encoding="utf-8"))["benchmark"]["root"]

    failures = 0

    for metrics_path in args.metrics:
        evaluations_dir = Path(args.output_dir) / "evaluations"

        if args.single_variant:
            record = score_run(metrics_path, root, evaluations_dir=evaluations_dir)
            detail = record.get("error") or ", ".join(
                f"{m['cell']} expected {m['expected']!r} got {m['actual']!r}"
                for m in record.get("mismatches", [])[:3]
            )
            summary = ""
        else:
            record = score_run_all_variants(metrics_path, root, evaluations_dir=evaluations_dir)
            summary = f"soft={record['soft']:.2f} hard={record['hard']} "
            summary += " ".join(
                f"v{v['variant']}:{'ok' if v['passed'] else 'FAIL'}" for v in record["variants"]
            )
            first_failure = next((v for v in record["variants"] if not v["passed"]), None)
            detail = ""

            if first_failure:
                reason = first_failure.get("error") or ", ".join(
                    f"{m['cell']} expected {m['expected']!r} got {m['actual']!r}"
                    for m in first_failure["mismatches"][:2]
                )
                detail = f"(v{first_failure['variant']}: {reason})"

            detail = record.get("error") or detail

        verdict = "PASS" if record["passed"] else "FAIL"
        runtime = record["runtime_status"]
        print(f"{verdict}  {record['task_id']} (runtime: {runtime})  {summary}{detail}")
        failures += not record["passed"]

    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
