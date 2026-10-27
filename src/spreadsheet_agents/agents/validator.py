"""Validate edited workbooks against task requirements.
"""

from pathlib import Path

from langchain_openai import ChatOpenAI
from openpyxl import load_workbook

from spreadsheet_agents.schemas.task import (
    EditPlan,
    TaskPlan,
    ValidationReport,
)
from spreadsheet_agents.tools.workbook_editor import (
    copy_cell_pairs,
    translate_copied_value,
)
from spreadsheet_agents.tools.workbook_reader import inspect_workbook

PROMPT_PATH = (
    Path(__file__).resolve().parents[1]
    / "prompts/baseline/validator.md"
)


def _written_cells(op) -> set[tuple[str, str]]:
    """Cells an operation writes (sheet, coordinate)."""

    if op.type in {"set_cell", "set_formula", "clear_cell"}:
        return {(op.sheet, op.cell)}

    if op.type == "copy_range":
        return {
            (op.destination_sheet, destination)
            for _, destination in copy_cell_pairs(
                op.source_range, op.destination_cell
            )
        }

    return set()


def _verify_copy(wb, operations, index) -> tuple[list[str], str | None]:
    """Check a copy_range against the output workbook.

    Returns (errors, reason_unverified). The check reads the source from
    the output, so it is only sound when nothing writes to the source or
    destination cells after the copy.
    """

    op = operations[index]

    try:
        pairs = copy_cell_pairs(op.source_range, op.destination_cell)
        later_writes = set().union(
            *(_written_cells(later) for later in operations[index + 1:])
        )
    except Exception as exc:
        return [], f"invalid range ({exc})"

    for sheet in (op.source_sheet, op.destination_sheet):
        if sheet not in wb.sheetnames:
            return [f"Missing sheet: {sheet}"], None

    touched = {(op.source_sheet, s) for s, _ in pairs} | {
        (op.destination_sheet, d) for _, d in pairs
    }

    if {(op.source_sheet, s) for s, _ in pairs} & {
        (op.destination_sheet, d) for _, d in pairs
    }:
        return [], "source and destination overlap"

    if touched & later_writes:
        return [], "a later operation changes the copied cells"

    errors = []

    for source, destination in pairs:
        expected = translate_copied_value(
            wb[op.source_sheet][source].value, source, destination
        )
        actual = wb[op.destination_sheet][destination].value

        if actual != expected:
            errors.append(
                f"{op.destination_sheet}!{destination}: expected "
                f"{expected!r} (copy of {op.source_sheet}!{source}), "
                f"got {actual!r}"
            )

    return errors, None


def deterministic_validate(
    output_path: str,
    edit_plan: EditPlan,
) -> ValidationReport:

    errors = []
    warnings = []
    checked = 0

    try:
        wb = load_workbook(output_path, data_only=False)
    except Exception as exc:
        return ValidationReport(
            passed=False,
            errors=[f"Output workbook is not readable: {exc}"],
        )

    # Operations run in order, so only the last write to a cell
    # decides what the output should contain.
    last_cell_op = {}

    try:
        for index, op in enumerate(edit_plan.operations):

            if op.type == "create_sheet":
                checked += 1

                if op.sheet not in wb.sheetnames:
                    errors.append(
                        f"Missing created sheet: {op.sheet}"
                    )

            elif op.type in {
                "set_cell",
                "set_formula",
                "clear_cell",
            }:
                last_cell_op[(op.sheet, op.cell)] = op

            elif op.type == "copy_range":
                copy_errors, unverified = _verify_copy(
                    wb, edit_plan.operations, index
                )

                if unverified:
                    warnings.append(
                        "copy_range was not verified "
                        f"({unverified}): "
                        f"{op.source_sheet}!{op.source_range} -> "
                        f"{op.destination_sheet}!{op.destination_cell}"
                    )
                else:
                    checked += 1
                    errors.extend(copy_errors)

        for op in last_cell_op.values():
            checked += 1

            if op.sheet not in wb.sheetnames:
                errors.append(
                    f"Missing sheet: {op.sheet}"
                )
                continue

            actual = wb[op.sheet][op.cell].value

            expected = (
                None
                if op.type == "clear_cell"
                else op.value
            )

            if actual != expected:
                errors.append(
                    f"{op.sheet}!{op.cell}: "
                    f"expected {expected!r}, got {actual!r}"
                )

    finally:
        wb.close()

    return ValidationReport(
        passed=len(errors) == 0,
        errors=errors,
        warnings=warnings,
        checked_operations=checked,
    )


def run_validator(
    user_request: str,
    task_plan: TaskPlan,
    edit_plan: EditPlan,
    output_path: str,
    model: str,
) -> ValidationReport:

    deterministic = deterministic_validate(
        output_path,
        edit_plan,
    )

    if not deterministic.passed:
        return deterministic

    output_info = inspect_workbook(output_path)

    llm = ChatOpenAI(
        model=model,
        temperature=0,
    ).with_structured_output(
        ValidationReport,
        method="function_calling",
    )

    prompt = PROMPT_PATH.read_text(encoding="utf-8")

    semantic = llm.invoke([
        ("system", prompt),
        (
            "human",
            f"User request:\n{user_request}\n\n"
            f"Task plan:\n"
            f"{task_plan.model_dump_json(indent=2)}\n\n"
            f"Edit plan:\n"
            f"{edit_plan.model_dump_json(indent=2)}\n\n"
            f"Output workbook preview:\n"
            f"{output_info.model_dump_json(indent=2)}"
        ),
    ])

    return ValidationReport(
        passed=deterministic.passed and semantic.passed,
        errors=deterministic.errors + semantic.errors,
        warnings=deterministic.warnings + semantic.warnings,
        checked_operations=deterministic.checked_operations,
    )