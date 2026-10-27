"""Read workbook contents and structural metadata.

"""
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from spreadsheet_agents.schemas.workbook import (
    SheetInfo,
    WorkbookInfo,
)


def _plain(value):
    """Turn openpyxl's special formula objects into JSON-friendly values."""

    if isinstance(value, ArrayFormula):
        return value.text

    if isinstance(value, DataTableFormula):
        return f"{{data table: {value.ref}}}"

    return value


def inspect_workbook(
    path: str | Path,
    max_preview_rows: int = 5,
    max_preview_columns: int = 20,
    max_formulas: int = 30,
) -> WorkbookInfo:
    path = Path(path)

    if not path.exists():
        raise FileNotFoundError(path)

    wb = load_workbook(path, read_only=True, data_only=False)

    sheets = []

    try:
        for ws in wb.worksheets:
            preview = []

            for row in ws.iter_rows(
                min_row=1,
                max_row=min(ws.max_row, max_preview_rows),
                max_col=min(ws.max_column, max_preview_columns),
                values_only=True,
            ):
                preview.append([_plain(value) for value in row])

            headers = (
                [
                    str(value) if value is not None else ""
                    for value in preview[0]
                ]
                if preview
                else []
            )

            formulas = []

            for row in ws.iter_rows(
                min_row=1,
                max_row=min(ws.max_row, 100),
                max_col=min(ws.max_column, 30),
            ):
                for cell in row:
                    if cell.data_type == "f":
                        formulas.append({
                            "cell": cell.coordinate,
                            "formula": str(_plain(cell.value)),
                        })

                        if len(formulas) >= max_formulas:
                            break

                if len(formulas) >= max_formulas:
                    break

            sheets.append(
                SheetInfo(
                    name=ws.title,
                    rows=ws.max_row,
                    columns=ws.max_column,
                    headers=headers,
                    sample_rows=preview,
                    formulas=formulas,
                )
            )

    finally:
        wb.close()

    return WorkbookInfo(
        path=str(path),
        sheets=sheets,
    )