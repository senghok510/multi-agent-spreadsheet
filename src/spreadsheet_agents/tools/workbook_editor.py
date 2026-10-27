"""Apply workbook edits while preserving source files.
"""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import (
    column_index_from_string,
    coordinate_from_string,
    get_column_letter,
    range_boundaries,
)

from spreadsheet_agents.schemas.task import (
    SpreadsheetOperation,
)


def _validate_cell(cell: str) -> None:
    coordinate_from_string(cell)


def refers_to_own_cell(formula: str, sheet: str, cell: str) -> bool:
    """True if `formula`, written into `sheet!cell`, reads that same cell.

    Covers direct references (=B2), absolute ones (=$B$2), and ranges or
    whole rows/columns that contain the cell (=SUM(B1:B5), =SUM(B:B)).
    References built at run time (INDIRECT, OFFSET) cannot be seen here.
    """

    letters, row = coordinate_from_string(cell)
    column = column_index_from_string(letters)

    for token in Tokenizer(formula).items:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue

        reference = token.value
        reference_sheet = None

        if "!" in reference:
            reference_sheet, reference = reference.rsplit("!", 1)
            reference_sheet = reference_sheet.strip("'").replace("''", "'")

        if reference_sheet is not None and reference_sheet != sheet:
            continue

        try:
            min_col, min_row, max_col, max_row = range_boundaries(
                reference.replace("$", "")
            )
        except ValueError:
            continue  # a defined name such as TaxRate, not a cell reference

        in_columns = min_col is None or min_col <= column <= max_col
        in_rows = min_row is None or min_row <= row <= max_row

        if in_columns and in_rows:
            return True

    return False


def translate_copied_value(value, source_cell: str, destination_cell: str):
    """Return what Excel would paste: formulas have relative references shifted."""

    if isinstance(value, str) and value.startswith("="):
        return Translator(value, origin=source_cell).translate_formula(
            destination_cell
        )

    return value


def copy_cell_pairs(
    source_range: str,
    destination_cell: str,
) -> list[tuple[str, str]]:
    """Pair every source coordinate with the coordinate it is copied to."""

    min_col, min_row, max_col, max_row = range_boundaries(source_range)
    dst_col, dst_row, _, _ = range_boundaries(destination_cell)

    return [
        (
            f"{get_column_letter(min_col + col_offset)}{min_row + row_offset}",
            f"{get_column_letter(dst_col + col_offset)}{dst_row + row_offset}",
        )
        for row_offset in range(max_row - min_row + 1)
        for col_offset in range(max_col - min_col + 1)
    ]


def _apply_operation(wb, operation: SpreadsheetOperation):
    op = operation

    if op.type == "create_sheet":
        if not op.sheet:
            raise ValueError("create_sheet requires sheet")

        if op.sheet in wb.sheetnames:
            raise ValueError(
                f"Sheet already exists: {op.sheet}"
            )

        wb.create_sheet(op.sheet)
        return

    if op.type in {"set_cell", "set_formula", "clear_cell"}:
        if not op.sheet or not op.cell:
            raise ValueError(
                f"{op.type} requires sheet and cell"
            )

        if op.sheet not in wb.sheetnames:
            raise ValueError(
                f"Unknown worksheet: {op.sheet}"
            )

        _validate_cell(op.cell)
        ws = wb[op.sheet]

        # set_cell with a string starting "=" is also stored as a formula.
        if (
            op.type in {"set_cell", "set_formula"}
            and isinstance(op.value, str)
            and op.value.startswith("=")
            and refers_to_own_cell(op.value, op.sheet, op.cell)
        ):
            raise ValueError(
                f"Formula in {op.sheet}!{op.cell} refers to its own "
                "cell (circular reference)"
            )

        if op.type == "clear_cell":
            ws[op.cell] = None

        elif op.type == "set_formula":
            if not isinstance(op.value, str):
                raise ValueError("Formula must be a string")

            if not op.value.startswith("="):
                raise ValueError(
                    "Excel formula must start with '='"
                )

            ws[op.cell] = op.value

        else:
            ws[op.cell] = op.value

        return

    if op.type == "copy_range":
        required = [
            op.source_sheet,
            op.source_range,
            op.destination_sheet,
            op.destination_cell,
        ]

        if any(value is None for value in required):
            raise ValueError(
                "copy_range requires source and destination"
            )

        if op.source_sheet not in wb.sheetnames:
            raise ValueError("Unknown source sheet")

        if op.destination_sheet not in wb.sheetnames:
            raise ValueError("Unknown destination sheet")

        src = wb[op.source_sheet]
        dst = wb[op.destination_sheet]

        # Snapshot first to support overlapping ranges.
        copies = [
            (
                destination,
                translate_copied_value(
                    src[source].value, source, destination
                ),
            )
            for source, destination in copy_cell_pairs(
                op.source_range, op.destination_cell
            )
        ]

        for destination, value in copies:
            dst[destination] = value

        return

    raise ValueError(
        f"Unsupported operation: {op.type}"
    )


def apply_operations(
    input_path: str | Path,
    output_path: str | Path,
    operations: list[SpreadsheetOperation],
    max_operations: int = 1000,
) -> Path:
    input_path = Path(input_path)
    output_path = Path(output_path)

    if len(operations) > max_operations:
        raise ValueError(
            f"Too many operations: {len(operations)}"
        )

    if not input_path.exists():
        raise FileNotFoundError(input_path)

    if output_path.resolve() == input_path.resolve():
        raise ValueError(
            "Refusing to overwrite the input workbook; "
            "choose a different output path"
        )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    wb = load_workbook(input_path)

    try:
        for index, operation in enumerate(operations):
            try:
                _apply_operation(wb, operation)
            except Exception as exc:
                raise RuntimeError(
                    f"Operation {index} failed: {exc}"
                ) from exc

        wb.save(output_path)

    finally:
        wb.close()

    return output_path